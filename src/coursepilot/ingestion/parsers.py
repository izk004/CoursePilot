from pathlib import Path
import fitz
from pptx import Presentation
from coursepilot.exceptions import ParsingError
from coursepilot.models import ParsedSection


def parse_file(filename: str, data: bytes) -> list[ParsedSection]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        try:
            document = fitz.open(stream=data, filetype="pdf")
            sections = [ParsedSection(title=f"第 {i} 页", text=page.get_text("text").strip(), location=i)
                        for i, page in enumerate(document, 1) if page.get_text("text").strip()]
            document.close()
        except Exception as exc:
            raise ParsingError(f"无法解析 PDF：{exc}") from exc
    elif suffix == ".pptx":
        try:
            import io
            deck = Presentation(io.BytesIO(data))
            sections = []
            for i, slide in enumerate(deck.slides, 1):
                text = "\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip())
                if text.strip(): sections.append(ParsedSection(title=f"第 {i} 张幻灯片", text=text.strip(), location=i))
        except Exception as exc:
            raise ParsingError(f"无法解析 PPTX：{exc}") from exc
    elif suffix == ".txt":
        text = ""
        for encoding in ("utf-8-sig", "utf-8", "gb18030"):
            try:
                text = data.decode(encoding); break
            except UnicodeDecodeError:
                continue
        if not text: raise ParsingError("TXT 必须是 UTF-8 或 GB18030 编码。")
        sections = [ParsedSection(title="文本资料", text=text.strip(), location=1)]
    else:
        raise ParsingError("仅支持 PDF、PPTX 和 TXT 文件。")
    if not sections: raise ParsingError("未能从资料中提取可用文本；扫描版 PDF 需要先 OCR。")
    return sections
